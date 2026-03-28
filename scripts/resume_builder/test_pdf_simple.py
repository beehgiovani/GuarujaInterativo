from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF
import os
import unicodedata

# --- DADOS ---
NAME = "Bruno Giovani Pereira"
LOCATION = "Guaruja, SP"
CONTACTS = "brunogp.corretor@gmail.com | (13) 99124-8146"
LINKEDIN = "https://www.linkedin.com/in/beehgiovani/"
GITHUB = "https://github.com/beehgiovani/"

PRIMARY_COLOR_RGB = (0, 74, 153)
PRIMARY_COLOR_HEX = RGBColor(0, 74, 153)

SUMMARY = "Software Engineer Full-Stack com especializacao em Geotecnologia (GIS) e Automacao Avancada. Certificacoes em Angular e Power BI."

# Simplificado para teste
EXPERIENCES = [{"title": "Senior Consultant", "period": "2024", "description": "Dev Kotlin", "bullets": ["Ajustes."]}]
CERTIFICATIONS = [{"nome": "Angular", "org": "Loiane", "date": "2026"}]
EDUCATION = ["Farmacia"]
SKILLS = {"Linguagens": "JS, Python"}

def remove_accents(input_str):
    if not isinstance(input_str, str): return input_str
    return "".join([c for c in unicodedata.normalize('NFKD', input_str) if not unicodedata.combining(c)])

class PDF(FPDF):
    def header(self): pass
    def add_top_header(self):
        print("Header start")
        self.set_font('helvetica', 'B', 22)
        self.set_text_color(*PRIMARY_COLOR_RGB)
        self.cell(190, 15, remove_accents(NAME), 0, 1, 'C')
        self.set_font('helvetica', '', 10)
        self.set_text_color(100, 100, 100)
        self.cell(190, 7, f"{remove_accents(LOCATION)} | {remove_accents(CONTACTS)}", 0, 1, 'C')
        self.cell(190, 7, f"LinkedIn: {LINKEDIN}", 0, 1, 'C')
        self.ln(5)
        print("Header end")

    def section_title(self, title):
        print(f"Section: {title}")
        self.set_font('helvetica', 'B', 12)
        self.set_fill_color(245, 248, 255) 
        self.set_text_color(*PRIMARY_COLOR_RGB)
        self.cell(190, 10, f"  {remove_accents(title)}", 0, 1, 'L', 1)
        self.ln(2)

    def body_text(self, text, bullet=False):
        self.set_font('helvetica', '', 10)
        self.set_text_color(50, 50, 50)
        prefix = "- " if bullet else ""
        self.multi_cell(190, 6, f"{prefix}{remove_accents(text)}")
        self.ln(1)

def create_pdf_final(path):
    print(f"Generating PDF: {path}")
    pdf = PDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.add_top_header()
    pdf.section_title('RESUMO')
    pdf.body_text(SUMMARY)
    pdf.section_title('CONTATOS')
    pdf.body_text(GITHUB, bullet=True)
    pdf.output(path)
    print("PDF Done")

if __name__ == "__main__":
    downloads_path = r"C:\Users\bruno\Downloads"
    output_pdf = os.path.join(downloads_path, "Curriculo_Teste_Simples.pdf")
    create_pdf_final(output_pdf)
